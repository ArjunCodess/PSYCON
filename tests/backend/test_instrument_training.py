"""Synthetic study fixtures live only in disposable PostgreSQL schemas."""
import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timezone
import pytest
import psycopg
from psycopg.types.json import Jsonb
from tests.backend.test_instrument import service,add
from backend.instrument import answers,datasets,training,jobs
from backend.instrument.store import uid,StaleJob


def setup_participant(service,sid,code='P01'):
    store=service.store
    from backend.instrument.provenance import review_source
    if not store.rows('SELECT id FROM source_reviews WHERE session_id=%s',(sid,)):
        review_source(store,sid,dict(kind='independent',reviewer_id='test-source-reviewer',reason='Synthetic fixture generated independently with a distinct phase'))
    speaker=store.one('SELECT * FROM speakers WHERE session_id=%s ORDER BY label LIMIT 1',(sid,))
    p=answers.add_participant(store,sid,dict(code=code))
    answers.map_participant(store,sid,p['id'],dict(speaker_id=speaker['id'],status='confirmed',reviewer_id='mapper',reason='Synthetic test mapping'))
    store.insert('consent_records',dict(id=uid(),participant_id=p['id'],status='documented',training_allowed=True,recorded_by='test-operator',source='Synthetic test consent'))
    return p,speaker


def upload(store,sid,pid,score=0,form='at-4.0',target='I',client=None):
    import csv
    text=io.StringIO(newline='');writer=csv.DictWriter(text,fieldnames=['item','score','fair_opportunity','confidence','evidence','validity']);writer.writeheader()
    writer.writerow(dict(item=target,score=score,fair_opportunity='yes',confidence=3,evidence=json.dumps([dict(kind='support',start_s=.2,end_s=1.,context='Synthetic observed exchange'),dict(kind='support',start_s=1.,end_s=2.,context='Second synthetic exchange')]),validity=json.dumps(dict(speaker_identity=dict(category='Yes'),recording_support=dict(category='Yes')))))
    options=dict(participant_id=pid,reviewer_id='rater',form_id=form)
    if client:
        from backend.instrument.import_jobs import process_import
        headers={'X-PSYCON-Request':'research-instrument'}
        response=client.post(f'/api/instrument/sessions/{sid}/answers/preview',headers=headers,data=dict(answers=(io.BytesIO(text.getvalue().encode()),'test-answers.csv'),options=json.dumps(options)))
        assert response.status_code==200,response.json
        preview=response.json
        assert not preview['errors']
        response=client.post(f'/api/instrument/sessions/{sid}/answers/commit',headers=headers,json=dict(import_id=preview['import_id'],durable=True))
        assert response.status_code==202,response.json
        job_id=response.json['id']
        process_import(store)
        job=client.get(f'/api/instrument/jobs/{job_id}').json
        assert job['status']=='complete',job
        saved=job['result']
        response=client.post('/api/instrument/submissions/'+saved['submissions'][0]+'/review',headers=headers,json=dict(reviewer_id='independent',decision='approved',notes='Synthetic independent review'))
        assert response.status_code==201,response.json
        return preview,saved
    preview=answers.preview(store,sid,text.getvalue().encode(),'test-answers.csv',options)
    assert not preview['errors']
    saved=answers.commit(store,sid,preview['import_id'])
    sub=saved['submissions'][0]
    store.insert('annotation_reviews',dict(id=uid(),submission_id=sub,reviewer_id='independent',decision='approved',notes='Synthetic independent review'))
    return preview,saved


def test_individual_import_states_duplicates_atomic_revisions(service):
    sid=add(service);p,s=setup_participant(service,sid)
    data=b'item,score,fair_opportunity,confidence\nA,0,yes,3\nB,N/O,no,1\nC,,,\n'
    options=dict(participant_id=p['id'],reviewer_id='rater')
    preview=answers.preview(service.store,sid,data,'sheet.csv',options)
    assert [r['state'] for r in preview['rows']]==['scored','not_observed','missing']
    saved=answers.commit(service.store,sid,preview['import_id']);assert len(saved['submissions'])==1
    duplicate=answers.preview(service.store,sid,data,'sheet.csv',options);assert duplicate['duplicate']
    assert answers.commit(service.store,sid,duplicate['import_id'])['duplicate']
    bad=answers.preview(service.store,sid,b'item,score\nA,7\nB,1\n','bad.csv',options)
    with pytest.raises(ValueError,match='Resolve every'): answers.commit(service.store,sid,bad['import_id'])
    assert len(service.store.rows('SELECT * FROM annotation_submissions'))==1
    correction=answers.preview(service.store,sid,b'item,score\nA,2\n','correction.csv',options)
    assert correction['rows'][0]['fair_opportunity'] is True and correction['rows'][0]['confidence']==3
    with pytest.raises(ValueError,match='reason'): answers.commit(service.store,sid,correction['import_id'])
    answers.commit(service.store,sid,correction['import_id'],'Corrected after source review')
    assert len(service.store.rows('SELECT * FROM annotation_answers'))==6  # B/C inherited.
    with pytest.raises(psycopg.Error): service.store.execute('UPDATE annotation_answers SET score=1 WHERE state=%s',('scored',))


def test_xlsx_formula_sheet_limits_and_batch_identity(service):
    from openpyxl import Workbook
    sid=add(service);p,s=setup_participant(service,sid)
    wb=Workbook();ws=wb.active;ws.append(['participant','item','score']);ws.append(['P01','A',0]);wb.create_sheet('Second')
    data=io.BytesIO();wb.save(data)
    assert answers.parse_file(data.getvalue(),'answers.xlsx')['needs_worksheet']
    parsed=answers.parse_file(data.getvalue(),'answers.xlsx',ws.title)
    assert not answers.normalize(service.store,sid,parsed['rows'],dict(participants={'P01':p['id']}))['errors']
    assert answers.normalize(service.store,sid,parsed['rows'],dict(participants={}))['errors']
    ws['C2']='=1+1';data=io.BytesIO();wb.save(data)
    with pytest.raises(ValueError,match='Formula'): answers.parse_file(data.getvalue(),'answers.xlsx',ws.title)
    preview=answers.preview(service.store,sid,b'participant,item,score\nP02,A,1\n','mismatch.csv',dict(participant_id=p['id']))
    assert preview['errors']


def test_concurrent_claims_expiry_cancellation_stale_completion(service):
    store=service.store
    for i in range(4):jobs.enqueue(store,'import','test-'+str(i),1)
    with ThreadPoolExecutor(max_workers=4) as executor:
        claimed=list(executor.map(lambda i:jobs.claim(store,'import',str(i)),range(4)))
    assert len({j['id'] for j in claimed})==4
    job=claimed[0];jobs.cancel(store,job['id'])
    with pytest.raises(StaleJob):jobs.finish(store,job,{'unsafe':'stale'})
    with pytest.raises(StaleJob):
        with jobs.lease(store,job):store.execute('DELETE FROM workspace_settings')
    expired=claimed[1]
    store.execute("UPDATE jobs SET lease_until=now()-interval '1 second' WHERE id=%s",(expired['id'],))
    store.execute("UPDATE worker_heartbeats SET seen_at=now()-interval '5 minutes' WHERE owner=%s",(expired['owner'],))
    assert jobs.claim(store,'import','replacement')['id']==expired['id']
    with pytest.raises(StaleJob):jobs.finish(store,expired,{})


def test_connected_person_duplicate_sources_and_reference_isolation():
    sessions=[dict(id='a',sha256='one',versions={}),dict(id='b',sha256='two',versions={}),dict(id='c',sha256='two',versions={})]
    people=[dict(person_id='person',session_id='a'),dict(person_id='person',session_id='b')]
    components=datasets.connected_components(sessions,people)
    assert len(set(components.values()))==1


def test_readiness_without_consent_and_mapping_is_explicit(service):
    sid=add(service);p=answers.add_participant(service.store,sid,dict(code='P01'))
    upload(service.store,sid,p['id'])
    ready=datasets.readiness(service.store,'at')
    assert not ready['rows'][0]['eligible']
    assert 'Training consent unavailable' in ready['rows'][0]['reasons']
    assert 'Needs reviewed speaker mapping' in ready['rows'][0]['reasons']
    snapshot=datasets.freeze(service.store,'at')
    run=training.queue(service.store,snapshot['id'])
    completed=training.process_training(service.store)
    assert completed['status']=='insufficient_data'
    assert not service.store.rows('SELECT * FROM model_versions')


def test_real_fit_portable_loading_activation_rollback_withdrawal(service):
    store=service.store;participants=[]
    from backend.instrument.app import create_app
    from backend.instrument.store import Store
    from backend.instrument.worker import run_trainer
    client=create_app(instrument=service,testing=True).test_client()
    headers={'X-PSYCON-Request':'research-instrument'}
    for day in range(1,19):
        role='development' if day<=12 else 'validation' if day<=15 else 'evaluation'
        sid=add(service,day=day,phase=day/11,split=role);p,s=setup_participant(service,sid)
        label=0 if day%2 else 4
        store.execute('UPDATE features SET value=%s WHERE speaker_id=%s AND name=%s',(.05 if label==0 else .95,s['id'],'question_ratio'))
        upload(store,sid,p['id'],label,client=client)
        upload(store,sid,p['id'],label,form='communication-1',target='questioning',client=client)
        participants.append((sid,p,s))
    for family in ('at','communication'):
        response=client.post('/api/instrument/training/snapshots',headers=headers,json=dict(family=family,seed=42,study='unseen_participant'))
        assert response.status_code==201,response.json
        snapshot=response.json
        summary=next(s['readiness'] for s in client.get('/api/instrument/training').json['snapshots'] if s['id']==snapshot['id'])
        assert summary['can_fit']
        assert summary['targets']['I' if family=='at' else 'questioning']['splits']['train']['records']==12
        assert snapshot['manifest']['counts']['examples']==18
        assert {e['split'] for e in snapshot['manifest']['examples']}=={'train','validation','test'}
        with pytest.raises(psycopg.Error):store.execute('UPDATE training_snapshots SET seed=1 WHERE id=%s',(snapshot['id'],))
        response=client.post('/api/instrument/training/runs',headers=headers,json=dict(snapshot_id=snapshot['id']))
        assert response.status_code==201,response.json
        run=response.json
        trainer=Store(store.root,database_url=store.url,schema=store.schema)
        completed=run_trainer(trainer,once=True)
        assert completed['status']=='complete',completed.get('error')
        assert completed['versions']['enqueue_runtime']==run['versions']
        assert completed['versions']['execution_runtime']==training.runtime_versions()

        model=store.one('SELECT * FROM model_versions WHERE run_id=%s',(run['id'],))
        artifact=training.load(store,model);assert artifact['features']
        assert model['manifest']['versions']==completed['versions']
        from backend.instrument.app import create_app
        client=create_app(instrument=service,testing=True).test_client()
        assert client.get('/api/instrument/training/snapshots/'+snapshot['id']+'/export').json['sha256']==snapshot['sha256']
        assert client.get('/api/instrument/models/'+model['id']+'/artifact').json==artifact
        with pytest.raises(ValueError,match='previously deployed'):
            training.activate(store,model['id'],'operator','Cannot call a new candidate a rollback',exploratory=True,rollback=True)
        training.activate(store,model['id'],'operator','Synthetic test activation',exploratory=True)
        result=training.predict(store,participants[-1][2]['id']);assert result['predictions']
        assert result['predictions'][-1]['score'] in (0,4)
        assert result['predictions'][-1]['family']==family
        assert result['predictions'][-1]['lineage']['snapshot_hash']==snapshot['sha256']
        from backend.instrument.research import packets
        packet=packets(service,participants[-1][2]['id'])
        assert packet['B']['supervised_predictions']==packet['C']['supervised_predictions']
        assert 'supervised_predictions' not in packet['A']
        predicted=next(p for p in packet['B']['supervised_predictions'] if p['model_id']==model['id'])
        assert predicted['lineage']['feature_schema']['version']=='conversation-2'
        assert predicted['lineage']['snapshot_hash']==snapshot['sha256']
        assert predicted['lineage']['rubric_source_hash']
        second=training.queue(store,snapshot['id']);training.process_training(store)
        candidate=store.one('SELECT * FROM model_versions WHERE run_id=%s',(second['id'],))
        training.activate(store,candidate['id'],'operator','Second synthetic version',exploratory=True)
        training.activate(store,model['id'],'operator','Synthetic rollback',exploratory=True,rollback=True)
    sid,p,s=participants[0]
    with store.connect() as db:answers.invalidate_participant(store,p['id'],'Synthetic withdrawal',db)
    assert all(m['status']=='needs_retraining' for m in store.rows('SELECT * FROM model_versions'))
    assert store.rows('SELECT * FROM snapshot_invalidations')
    assert client.get('/api/instrument/models/'+model['id']+'/artifact').status_code==400
    with pytest.raises(ValueError,match='withdrawn or corrected'):training.activate(store,model['id'],'operator','Must refuse',exploratory=True)


def test_backup_restore_assets_and_deletion_remove_frozen_copies(service,tmp_path,tmp_path_factory):
    from backend.instrument.backup import coordinated_backup,restore_backup
    from backend.instrument.deletion import delete_participant
    from psycopg import sql
    from urllib.parse import urlsplit
    from backend.instrument.assets import resolve
    sid=add(service);p,speaker=setup_participant(service,sid)
    upload(service.store,sid,p['id'])
    snapshot=datasets.freeze(service.store,'at')
    backup_directory=tmp_path_factory.mktemp('verified-backup')
    backup=coordinated_backup(service.store,backup_directory)
    assert backup['database']['restoration_verified']
    name='psycon_restore_'+uid().replace('-','')
    url=service.store.url
    with psycopg.connect(url,autocommit=True) as db:db.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(name)))
    destination_url=urlsplit(url)._replace(path='/'+name).geturl()
    try:
        restored=restore_backup(destination_url,backup_directory,backup_directory/'restored-assets')
        assert restored['database_restored'] and restored['asset_count']==2
        assert {a['original_filename'] for a in backup['assets']}=={'conversation.wav','normalized.wav'}
        with psycopg.connect(destination_url) as db:
            db.execute(sql.SQL('SET search_path TO {}').format(sql.Identifier(service.store.schema)))
            assert db.execute('SELECT count(*) FROM annotation_answers').fetchone()[0]==1
    finally:
        with psycopg.connect(url,autocommit=True) as db:db.execute(sql.SQL('DROP DATABASE {} WITH (FORCE)').format(sql.Identifier(name)))
    delete_participant(service.store,sid,p['id'])
    assert not service.store.rows('SELECT * FROM training_snapshots')
    assert not service.store.rows('SELECT * FROM annotation_answers')
    assert service.store.one('SELECT source_bytes FROM annotation_imports')['source_bytes'] is None
    assert service.store.rows('SELECT * FROM deletion_tombstones')


def test_retry_runtime_mismatch_is_a_durable_failure(service):
    snapshot=datasets.freeze(service.store,'at')
    run=training.queue(service.store,snapshot['id'])
    versions=dict(run['versions'],execution_runtime=dict(code_hash='different-runtime'))
    service.store.execute('UPDATE training_runs SET versions=%s WHERE id=%s',(Jsonb(versions),run['id']))
    completed=training.process_training(service.store)
    assert completed['status']=='failed'
    assert 'Retry runtime differs' in completed['error']
    assert service.store.one('SELECT status FROM jobs WHERE id=%s',(run['job_id'],))['status']=='failed'
    assert not service.store.rows('SELECT id FROM model_versions')


def test_training_retry_remains_queued_and_exhaustion_is_visible(service,monkeypatch):
    snapshot=datasets.freeze(service.store,'at')
    run=training.queue(service.store,snapshot['id'])
    def unavailable(*args):
        raise ConnectionError('Synthetic temporary outage')
    monkeypatch.setattr(training,'fit',unavailable)
    completed=training.process_training(service.store)
    assert completed['status']=='queued'
    job=service.store.one('SELECT * FROM jobs WHERE id=%s',(run['job_id'],))
    assert job['status']=='queued' and job['attempt']==1
    service.store.execute("UPDATE jobs SET status='failed',error='Retry budget exhausted' WHERE id=%s",(run['job_id'],))
    assert training.process_training(service.store) is None
    assert service.store.one('SELECT status,error FROM training_runs WHERE id=%s',(run['id'],))==dict(status='failed',error='Retry budget exhausted')


def test_idle_queue_poll_does_not_accumulate_worker_records(service):
    store=service.store
    before=store.one('SELECT count(*) AS n FROM worker_heartbeats')['n']
    for _ in range(3):
        assert jobs.claim(store,'training') is None
    assert store.one('SELECT count(*) AS n FROM worker_heartbeats')['n']==before
