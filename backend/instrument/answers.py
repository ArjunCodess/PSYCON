"""Bounded spreadsheet sources, guided previews, and immutable human revisions."""
import csv
import hashlib
import io
import json
import math
from pathlib import PurePath
import zipfile
from psycopg.types.json import Jsonb
from .store import uid, decode, public
from .forms import FLAGS, VALIDITY, PATTERNS

MAX_BYTES=10*1024*1024
MAX_ROWS=20000
COLUMNS={'participant','item','score','fair_opportunity','confidence','missing_reason','narrative','component','evidence','metadata','context_flags','validity','patterns','reviewer_id','source_type'}
METADATA_FIELDS={'session_id','date','participant_code','seat_channel','observer_code','class_section','topic','observation_minutes','languages','recording_quality','contextual_account','evidence_summary','signature','signed_at','reviewer_code','reviewed_at'}


def participant(store, sid, pid):
    return store.one('SELECT * FROM session_participants WHERE id=%s AND session_id=%s AND owner_id=%s', (pid,sid,'local'))


def add_participant(store,sid,body):
    store.one("SELECT id FROM sessions WHERE id=%s AND owner_id='local'", (sid,))
    code=str(body.get('code','')).strip()
    if not code or len(code)>100:
        raise ValueError('Use an anonymous participant code of 1–100 characters')
    person=body.get('person_id') or None
    if person:
        store.one("SELECT id FROM people WHERE id=%s AND owner_id='local' AND withdrawn_at IS NULL",(person,))
    row=dict(id=uid(),session_id=sid,person_id=person,code=code,owner_id='local')
    store.insert('session_participants',row)
    return participant(store,sid,row['id'])


def invalidate_participant(store,pid,reason,db):
    snapshots=db.execute('SELECT DISTINCT snapshot_id FROM training_examples WHERE participant_id=%s',(pid,)).fetchall()
    for row in snapshots:
        db.execute('INSERT INTO snapshot_invalidations VALUES (%s,%s,%s,now())',(uid(),row['snapshot_id'],reason))
        db.execute("UPDATE jobs SET status='canceled',error=%s WHERE subject_id IN (SELECT id::text FROM training_runs WHERE snapshot_id=%s) AND status IN ('queued','running')",(reason,row['snapshot_id']))
        db.execute("UPDATE training_runs SET status='canceled',error=%s WHERE snapshot_id=%s AND status IN ('queued','training')",(reason,row['snapshot_id']))
        db.execute("UPDATE model_versions SET status='needs_retraining' WHERE run_id IN (SELECT id FROM training_runs WHERE snapshot_id=%s)",(row['snapshot_id'],))
    sid=db.execute('SELECT session_id FROM session_participants WHERE id=%s',(pid,)).fetchone()['session_id']
    db.execute('UPDATE sessions SET input_revision=input_revision+1 WHERE id=%s',(sid,))
    db.execute("UPDATE jobs SET status='canceled',error=%s WHERE subject_id=%s AND status IN ('queued','running')",(reason,sid))
    db.execute('DELETE FROM model_predictions WHERE session_id=%s',(sid,))
    db.execute("UPDATE llm_runs SET status='stale',error=%s WHERE session_id=%s AND status IN ('queued','running','complete')",(reason,sid))


def map_participant(store,sid,pid,body):
    participant(store,sid,pid)
    speaker=body.get('speaker_id') or None; status=body.get('status','uncertain')
    reviewer=str(body.get('reviewer_id','')).strip()
    if status not in ('confirmed','uncertain','unmapped') or not reviewer:
        raise ValueError('A reviewer and mapping status are required')
    if speaker:
        store.one('SELECT id FROM speakers WHERE id=%s AND session_id=%s',(speaker,sid))
    if status=='confirmed' and not speaker:
        raise ValueError('A confirmed mapping requires a speaker')
    with store.connect() as db:
        db.execute('SELECT id FROM sessions WHERE id=%s FOR UPDATE',(sid,))
        if speaker and status=='confirmed':
            conflicts=db.execute("SELECT m.participant_id FROM speaker_mappings m WHERE m.session_id=%s AND m.speaker_id=%s AND m.status='confirmed' AND m.participant_id<>%s AND m.revision=(SELECT max(revision) FROM speaker_mappings x WHERE x.participant_id=m.participant_id)",(sid,speaker,pid)).fetchall()
            if conflicts:
                raise ValueError('This speaker already maps to another participant; review diarization first')
        revision=db.execute('SELECT coalesce(max(revision),0)+1 AS n FROM speaker_mappings WHERE participant_id=%s',(pid,)).fetchone()['n']
        previous=db.execute('SELECT speaker_id FROM speaker_mappings WHERE participant_id=%s ORDER BY revision DESC LIMIT 1',(pid,)).fetchone()
        person=db.execute('SELECT person_id FROM session_participants WHERE id=%s',(pid,)).fetchone()['person_id']
        if previous and previous['speaker_id'] and person:
            db.execute('UPDATE speakers SET profile_id=NULL WHERE id=%s AND profile_id=%s',(previous['speaker_id'],str(person)))
        row=dict(id=uid(),participant_id=pid,session_id=sid,speaker_id=speaker,revision=revision,status=status,reviewer_id=reviewer,reason=str(body.get('reason','')))
        store.insert('speaker_mappings',row,db)
        if status=='confirmed':
            if person:
                db.execute('UPDATE speakers SET profile_id=%s WHERE id=%s',(str(person),speaker))
        invalidate_participant(store,pid,'Participant mapping changed',db)
    store.invalidate()
    return row


def parse_file(data,filename,worksheet=None):
    if not data or len(data)>MAX_BYTES:
        raise ValueError('Annotation sources must contain 1 byte to 10 MiB')
    suffix=PurePath(filename).suffix.lower()
    if suffix=='.csv':
        try:
            text=data.decode('utf-8-sig'); reader=csv.DictReader(io.StringIO(text))
            headers=reader.fieldnames or []; rows=[]
            if len(headers)>150:
                raise ValueError('Too many CSV columns')
            for i,row in enumerate(reader):
                if i>=MAX_ROWS or None in row:
                    raise ValueError('Too many rows or a row has more cells than the header')
                rows.append(row)
        except UnicodeDecodeError as exc:
            raise ValueError('CSV must use UTF-8') from exc
        sheets=[]
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(f.file_size for f in z.infolist())>50*1024*1024 or len(z.infolist())>1000:
                raise ValueError('Expanded workbook exceeds the 50 MiB limit')
            if any('vbaProject' in f.filename for f in z.infolist()):
                raise ValueError('Macros are not supported')
        workbook=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        try:
            sheets=workbook.sheetnames
            if len(sheets)>1 and not worksheet:
                return dict(worksheets=sheets,headers=[],rows=[],needs_worksheet=True)
            if worksheet and worksheet not in sheets:
                raise ValueError('Select an existing worksheet')
            sheet=workbook[worksheet or sheets[0]]
            iterator=sheet.iter_rows(); first=next(iterator,None)
            headers=[str(c.value or '').strip() for c in (first or [])]; rows=[]
            if len(headers)>150:
                raise ValueError('Too many workbook columns')
            for i,cells in enumerate(iterator):
                if i>=MAX_ROWS:
                    raise ValueError('Workbook exceeds 20,000 rows')
                if any(c.data_type=='f' for c in cells):
                    raise ValueError(f'Formula in row {i+2}; upload literal answers')
                from datetime import date,datetime
                rows.append({headers[j]:c.value.isoformat() if isinstance(c.value,(date,datetime)) else c.value for j,c in enumerate(cells) if j<len(headers)})
        finally:
            workbook.close()
    else:
        raise ValueError('Upload ordinary CSV or XLSX; scanned forms are attachments')
    if not headers or any(not h for h in headers) or len(set(headers))!=len(headers):
        raise ValueError('Headers must be non-empty and unique')
    return dict(worksheets=sheets,headers=headers,rows=rows,needs_worksheet=False)


def boolean(value):
    if value in (None,''):
        return None
    if value is True or str(value).lower() in ('true','yes','1'):
        return True
    if value is False or str(value).lower() in ('false','no','0'):
        return False
    raise ValueError('Fair opportunity must be Yes, No, or empty')


def structured(value,default):
    if value in (None,''):
        return default
    try:
        result=decode(value)
    except (ValueError,TypeError) as exc:
        raise ValueError('Structured fields must contain JSON') from exc
    if not isinstance(result,type(default)):
        raise ValueError('Structured field has the wrong type')
    return result


def validate_metadata(metadata,sid,code):
    from datetime import datetime
    if metadata.get('session_id') and metadata['session_id']!=sid:
        raise ValueError('Form metadata names another session')
    if metadata.get('participant_code') and metadata['participant_code']!=code:
        raise ValueError('Form metadata names another participant')
    if 'recording_quality' in metadata and (type(metadata['recording_quality']) is not int or not 1<=metadata['recording_quality']<=5):
        raise ValueError('Recording quality must be an integer from 1 to 5')
    if 'observation_minutes' in metadata and (type(metadata['observation_minutes']) not in (int,float) or not math.isfinite(metadata['observation_minutes']) or metadata['observation_minutes']<0):
        raise ValueError('Observation minutes must be finite and non-negative')
    if 'languages' in metadata and (not isinstance(metadata['languages'],list) or any(not isinstance(v,str) for v in metadata['languages'])):
        raise ValueError('Languages must be an array of language codes or names')
    for field in ('signed_at','reviewed_at'):
        if metadata.get(field):
            try:
                value=datetime.fromisoformat(metadata[field].replace('Z','+00:00'))
                if value.tzinfo is None: raise ValueError('Timezone missing')
            except (ValueError,TypeError,AttributeError) as exc:
                raise ValueError(field+' requires an ISO timestamp with timezone') from exc
    if metadata.get('date'):
        from datetime import date
        try:date.fromisoformat(metadata['date'])
        except (ValueError,TypeError) as exc:raise ValueError('Form date must use YYYY-MM-DD') from exc
    return metadata


def normalize(store,sid,raw,options):
    form=store.one('SELECT * FROM form_definitions WHERE id=%s',(options.get('form_id','at-4.0'),))
    keys={r['key'] for r in store.rows('SELECT key FROM form_items WHERE form_id=%s',(form['id'],))}
    session=store.session(sid); selected=options.get('participant_id'); mappings=options.get('participants',{})
    columns=options.get('columns',{}); seen=set(); rows=[]; errors=[]
    for index,source in enumerate(raw):
        try:
            targets=[columns.get(k,k) for k in source if columns.get(k,k)!='ignore']
            if len(set(targets))!=len(targets): raise ValueError('Several columns map to the same field')
            row={columns.get(k,k):v for k,v in source.items() if columns.get(k,k)!='ignore'}
            unknown=set(row)-COLUMNS-keys-METADATA_FIELDS-set(FLAGS)-{'class'}
            if unknown:
                raise ValueError('Map or ignore unknown columns: '+', '.join(sorted(unknown)))
            code=str(row.get('participant') or '').strip()
            pid=selected or mappings.get(code)
            if not pid:
                raise ValueError('Explicit participant mapping is required')
            person=participant(store,sid,pid)
            if selected and code and code not in (person['code'],str(person['id'])):
                raise ValueError('Embedded participant conflicts with selected participant')
            if person['withdrawn_at']:
                raise ValueError('Participant has withdrawn')
            source_type=row.get('source_type') or options.get('source_type','observer')
            reviewer=str(row.get('reviewer_id') or options.get('reviewer_id') or '').strip() or None
            if source_type not in ('observer','self_report','independent_review'):
                raise ValueError('Unknown human source type')
            metadata=structured(row.get('metadata'),{})
            for field in METADATA_FIELDS:
                if row.get(field) in (None,''):continue
                value=row[field]
                if field=='recording_quality':
                    if str(value) not in ('1','2','3','4','5'):raise ValueError('Recording quality must be from 1 to 5')
                    value=int(value)
                elif field=='observation_minutes':value=float(value)
                elif field=='languages':value=[part.strip() for part in str(value).split(';') if part.strip()]
                elif field=='date':value=str(value).split('T')[0]
                if field in metadata and metadata[field]!=value:raise ValueError('Conflicting metadata field: '+field)
                metadata[field]=value
            if source_type!='self_report' and metadata.get('observer_code'):
                if reviewer and reviewer!=metadata['observer_code']:raise ValueError('Selected reviewer conflicts with the form observer code')
                reviewer=metadata['observer_code']
            if row.get('class') or row.get('class_section'):
                metadata['class_section']=row.get('class_section') or row.get('class')
            metadata=validate_metadata(metadata,sid,person['code'])
            flags=structured(row.get('context_flags'),{}); validity=structured(row.get('validity'),{}); patterns=structured(row.get('patterns'),{})
            for field in FLAGS:
                if row.get(field) not in (None,''): flags[field]=dict(present=boolean(row[field]),notes='')
            if set(flags)-set(FLAGS) or set(validity)-set(VALIDITY) or set(patterns)-set(PATTERNS):
                raise ValueError('Unknown context, validity, or pattern field')
            for value in flags.values():
                if boolean(value.get('present') if isinstance(value,dict) else value) is None: raise ValueError('Context flags require explicit Yes or No')
            for key,value in validity.items():
                if not isinstance(value,dict) or value.get('category') not in VALIDITY[key]:
                    raise ValueError('Invalid category for '+key)
            for value in patterns.values():
                if not isinstance(value,dict) or value.get('category') not in ('Observed','Not clear','N/O') or value.get('confidence') not in (1,2,3):
                    raise ValueError('Patterns require an exact category and confidence 1–3')
            answers=[(row.get('item'),row.get('score'))] if row.get('item') else [(k,row.get(k)) for k in sorted(keys)]
            for key,value in answers:
                if key not in keys or (pid,source_type,reviewer,key) in seen:
                    raise ValueError('Unknown or duplicate participant/item/reviewer')
                seen.add((pid,source_type,reviewer,key))
                state='missing' if value in (None,'') else 'not_observed' if str(value).strip().upper()=='N/O' else 'scored'
                score=None
                if state=='scored':
                    if isinstance(value,bool) or (type(value) in (int,float) and (not math.isfinite(value) or value not in range(5))) or (type(value) not in (int,float) and str(value).strip() not in ('0','1','2','3','4')):
                        raise ValueError('Scores must be integer 0–4, N/O, or empty')
                    score=int(value)
                confidence=row.get('confidence')
                if confidence not in (None,''):
                    if str(confidence) not in ('1','2','3'):
                        raise ValueError('Confidence must be 1–3')
                    confidence=int(confidence)
                else:
                    confidence=None
                evidence=structured(row.get('evidence'),[])
                for window in evidence:
                    a,b=window.get('start_s'),window.get('end_s')
                    if type(a) not in (int,float) or type(b) not in (int,float) or not math.isfinite(a+b) or not 0<=a<b or session['duration'] is None or b>session['duration']+.01:
                        raise ValueError('Evidence window must fit the processed recording')
                    if window.get('kind','support') not in ('support','event','response','same_session_baseline'):
                        raise ValueError('Invalid window kind')
                    if window.get('evidence_id'):
                        store.one('SELECT id FROM evidence WHERE id=%s AND session_id=%s',(window['evidence_id'],sid))
                current=store.rows('SELECT revision FROM annotation_submissions WHERE participant_id=%s AND form_id=%s AND source_type=%s AND reviewer_id IS NOT DISTINCT FROM %s ORDER BY revision DESC LIMIT 1',(pid,form['id'],source_type,reviewer))
                rows.append(dict(row_number=index+2,provided_fields=sorted(row),participant_id=pid,participant_code=person['code'],form_id=form['id'],source_type=source_type,reviewer_id=reviewer,expected_revision=current[0]['revision'] if current else 0,item_key=key,state=state,score=score,fair_opportunity=boolean(row.get('fair_opportunity')),confidence=confidence,missing_reason=str(row.get('missing_reason') or ''),narrative=str(row.get('narrative') or ''),component=str(row.get('component') or 'audio'),evidence=evidence,metadata=metadata,context_flags=flags,validity=validity,patterns=patterns))
        except (ValueError,LookupError) as exc:
            errors.append(dict(row=index+2,error=str(exc)))
    if not rows and not errors:
        errors.append(dict(row=0,error='No answers found'))
    shared={}
    for row in rows:
        group=shared.setdefault((row['participant_id'],row['source_type'],row['reviewer_id']),{})
        for field in ('metadata','context_flags','validity','patterns'):
            values=group.setdefault(field,{})
            for key,value in row[field].items():
                if key in values and values[key]!=value:errors.append(dict(row=row['row_number'],error='Conflicting form-level '+field+': '+key))
                values[key]=value
    for row in rows:
        for field,values in shared[(row['participant_id'],row['source_type'],row['reviewer_id'])].items():row[field]=values
    # Show exactly the fields a partial correction will retain before confirmation.
    for row in rows:
        previous=store.rows('SELECT * FROM annotation_submissions WHERE participant_id=%s AND form_id=%s AND source_type=%s AND reviewer_id IS NOT DISTINCT FROM %s ORDER BY revision DESC LIMIT 1',(row['participant_id'],row['form_id'],row['source_type'],row['reviewer_id']))
        if not previous:continue
        old=previous[0]
        row['metadata']={**old['metadata'],**row['metadata']}
        prior=store.rows('SELECT * FROM annotation_answers WHERE submission_id=%s AND item_key=%s',(old['id'],row['item_key']))
        if prior:
            for field in ('fair_opportunity','confidence','missing_reason','narrative','component'):
                if field not in row['provided_fields']:row[field]=prior[0][field]
            if 'evidence' not in row['provided_fields']:
                row['evidence']=store.rows('SELECT * FROM annotation_evidence WHERE answer_id=%s',(prior[0]['id'],))
        for table,field in (('annotation_context_flags','context_flags'),('annotation_validity_checks','validity'),('annotation_pattern_summaries','patterns')):
            prior_fields={r['name']:{k:v for k,v in r.items() if k not in ('submission_id','name')} for r in store.rows('SELECT * FROM '+table+' WHERE submission_id=%s',(old['id'],))}
            row[field]={**prior_fields,**row[field]}
    return dict(rows=rows,errors=errors,form_id=form['id'])


def preview(store,sid,data,filename,options):
    store.one("SELECT id FROM sessions WHERE id=%s AND owner_id='local'",(sid,))
    parsed=parse_file(data,filename,options.get('worksheet'))
    if parsed['needs_worksheet']:
        return parsed
    result=normalize(store,sid,parsed['rows'],options)
    result.update(headers=parsed['headers'],worksheets=parsed['worksheets'])
    digest=hashlib.sha256(data).hexdigest()
    existing=store.rows('SELECT id,preview,status FROM annotation_imports WHERE session_id=%s AND sha256=%s AND mapping=%s',(sid,digest,Jsonb(options)))
    if existing:
        return dict(**existing[0]['preview'],import_id=existing[0]['id'],duplicate=True,status=existing[0]['status'])
    key=uid()
    with store.connect() as db:
        store.insert('annotation_imports',dict(id=key,session_id=sid,owner_id='local',filename=filename,sha256=digest,source_bytes=data,media_type='text/csv' if filename.lower().endswith('.csv') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',mapping=options,preview=result),db)
        for i,row in enumerate(parsed['rows']):
            store.insert('annotation_import_rows',dict(import_id=key,row_number=i+2,data=row),db)
    return dict(**result,import_id=key,duplicate=False,status='preview')


def commit(store,sid,import_id,reason=''):
    with store.connect() as db:
        source=db.execute("SELECT * FROM annotation_imports WHERE id=%s AND session_id=%s AND owner_id='local' FOR UPDATE",(import_id,sid)).fetchone()
        if not source:
            raise LookupError('Import does not belong to this session')
        if source['status']=='committed':
            return dict(status='committed',duplicate=True)
        result=source['preview']
        if result['errors']:
            raise ValueError('Resolve every preview error before saving')
        groups={}
        for row in result['rows']:
            groups.setdefault((row['participant_id'],row['form_id'],row['source_type'],row['reviewer_id']),[]).append(row)
        saved=[]
        for (pid,form,kind,reviewer),rows in sorted(groups.items(),key=lambda x:str(x[0])):
            person=db.execute('SELECT * FROM session_participants WHERE id=%s AND session_id=%s FOR UPDATE',(pid,sid)).fetchone()
            if not person or person['withdrawn_at']:
                raise ValueError('Participant is absent or withdrawn')
            old=db.execute('SELECT * FROM annotation_submissions WHERE participant_id=%s AND form_id=%s AND source_type=%s AND reviewer_id IS NOT DISTINCT FROM %s ORDER BY revision DESC LIMIT 1',(pid,form,kind,reviewer)).fetchone()
            current=old['revision'] if old else 0
            if any(r['expected_revision']!=current for r in rows):
                raise ValueError('Answers changed after preview; create a fresh preview')
            if old and not reason.strip():
                raise ValueError('A correction requires a reason')
            submission=uid(); first=rows[0]
            first=dict(first,metadata={**(old['metadata'] if old else {}),**first['metadata']})
            store.insert('annotation_submissions',dict(id=submission,participant_id=pid,form_id=form,import_id=import_id,source_type=kind,reviewer_id=reviewer,uploader_id='local',revision=current+1,supersedes=str(old['id']) if old else None,correction_reason=reason,metadata=first['metadata'],raw=rows),db)
            # Partial corrections inherit unchanged prior answers without mutating their source revision.
            if old:
                submitted={r['item_key'] for r in rows}
                for previous in db.execute('SELECT * FROM annotation_answers WHERE submission_id=%s',(old['id'],)):
                    if previous['item_key'] in submitted:
                        continue
                    inherited_id=uid()
                    store.insert('annotation_answers',dict(id=inherited_id,submission_id=submission,**{k:previous[k] for k in ('item_key','state','score','fair_opportunity','confidence','missing_reason','narrative','component')}),db)
                    for prior in db.execute('SELECT * FROM annotation_evidence WHERE answer_id=%s',(previous['id'],)):
                        store.insert('annotation_evidence',dict(id=uid(),answer_id=inherited_id,**{k:prior[k] for k in ('evidence_id','kind','start_s','end_s','context','response','participation_effect')}),db)
            for row in rows:
                if old:
                    previous=db.execute('SELECT * FROM annotation_answers WHERE submission_id=%s AND item_key=%s',(old['id'],row['item_key'])).fetchone()
                    if previous:
                        for field in ('fair_opportunity','confidence','missing_reason','narrative','component'):
                            if field not in row.get('provided_fields',[]):row[field]=previous[field]
                        if 'evidence' not in row.get('provided_fields',[]):
                            row['evidence']=[public(dict(window)) for window in db.execute('SELECT * FROM annotation_evidence WHERE answer_id=%s',(previous['id'],))]
                aid=uid()
                store.insert('annotation_answers',dict(id=aid,submission_id=submission,**{k:row[k] for k in ('item_key','state','score','fair_opportunity','confidence','missing_reason','narrative','component')}),db)
                for window in row['evidence']:
                    store.insert('annotation_evidence',dict(id=uid(),answer_id=aid,evidence_id=window.get('evidence_id'),kind=window.get('kind','support'),start_s=window['start_s'],end_s=window['end_s'],context=str(window.get('context','')),response=str(window.get('response','')),participation_effect=str(window.get('participation_effect',''))),db)
            for name,value in first['context_flags'].items():
                store.insert('annotation_context_flags',dict(submission_id=submission,name=name,present=boolean(value.get('present')) if isinstance(value,dict) else boolean(value),notes=str(value.get('notes','')) if isinstance(value,dict) else ''),db)
            for name,value in first['validity'].items():
                store.insert('annotation_validity_checks',dict(submission_id=submission,name=name,category=value['category'],explanation=str(value.get('explanation',''))),db)
            for name,value in first['patterns'].items():
                store.insert('annotation_pattern_summaries',dict(submission_id=submission,name=name,category=value['category'],confidence=value['confidence'],narrative=str(value.get('narrative',''))),db)
            if old:
                for table,field in (('annotation_context_flags','context_flags'),('annotation_validity_checks','validity'),('annotation_pattern_summaries','patterns')):
                    for prior in db.execute('SELECT * FROM '+table+' WHERE submission_id=%s',(old['id'],)):
                        if prior['name'] not in first[field]:
                            store.insert(table,dict(prior,submission_id=submission),db)
            invalidate_participant(store,pid,'Human annotation revision changed',db)
            saved.append(submission)
        db.execute("UPDATE annotation_imports SET status='committed' WHERE id=%s",(import_id,))
    return dict(status='committed',submissions=saved,duplicate=False)
